"""
АкБарак — шаблоны документов (заявления, объяснительные, титульные листы).

Описание шаблонов лежит в templates.json — тот же файл читает сайт для предпросмотра,
поэтому лист на телефоне и напечатанный лист совпадают.
Здесь шаблон заполняется данными студента и собирается в DOCX (Times New Roman 14,
поля по ГОСТ). DOCX печатает программа на терминале через Word/LibreOffice.
"""

import json
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(HERE, "templates.json"), encoding="utf-8") as f:
    DATA = json.load(f)

FIELDS = DATA["fields"]
TEMPLATES = {t["id"]: t for t in DATA["templates"]}

FONT = "Times New Roman"
TOKEN = re.compile(r"\{(\w+)(?::([^{}|]*)\|([^{}]*))?\}")


class TemplateError(ValueError):
    pass


def clean_values(tpl, raw):
    """Проверяет поля шаблона. Возвращает словарь значений или бросает TemplateError."""
    values = {}
    for key in tpl["fields"]:
        spec = FIELDS[key]
        val = str(raw.get(key, "") or "").strip()
        val = re.sub(r"\s+", " ", val)[: spec.get("max", 120)]
        if spec.get("type") == "select":
            allowed = [o[0] for o in spec["options"]]
            if val not in allowed:
                val = spec.get("def", allowed[0])
        if not val and not spec.get("optional"):
            raise TemplateError(f"Заполните поле «{spec['label']}»")
        values[key] = val
    return values


def fill(text, values):
    """Подставляет значения. Возвращает (текст, есть_ли_пустые_поля)."""
    empty = False

    def sub(m):
        nonlocal empty
        key, male, female = m.group(1), m.group(2), m.group(3)
        if male is not None:  # {sex:вариант м|вариант ж}
            return female if values.get(key) == "f" else male
        val = values.get(key, "")
        if not val:
            empty = True
        return val

    return TOKEN.sub(sub, text), empty


def _font(run, size, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def build_docx(template_id, raw_values, path):
    tpl = TEMPLATES.get(template_id)
    if not tpl:
        raise TemplateError("Неизвестный шаблон")
    values = clean_values(tpl, raw_values)

    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(14)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    for b in tpl["blocks"]:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(b.get("before", 0))
        pf.space_after = Pt(b.get("after", 0))
        pf.line_spacing = 1.15

        if b.get("type") == "sign":
            # дата слева, подпись справа
            pf.tab_stops.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT)
            date, _ = fill("{date}", values)
            short, _ = fill("{fio_short}", values)
            _font(p.add_run(f"{date}\t____________  {short}"), 14)
            continue

        text, empty = fill(b["t"], values)
        if empty and b.get("opt"):
            p._element.getparent().remove(p._element)
            continue
        if b.get("u"):
            text = text.upper()

        align = b.get("a", "j")
        if align == "c":
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "hdr":
            pf.left_indent = Cm(8.5)
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        else:
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.line_spacing = 1.5
        if b.get("ind"):
            pf.first_line_indent = Cm(1.25)
        _font(p.add_run(text), b.get("s", 14), b.get("b", False))

    doc.save(path)
    return tpl["title"]
